from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
import re,subprocess,cairosvg,io
from PIL import Image

OUT=Path(__file__).resolve().parent
BASE=OUT.parent/"5g-thesis-presentation-5min"/"tesi-5min.pptx"
def X(p):return Inches(p/96)
def C(v):return RGBColor.from_string(v)
def box(sl,x,y,w,h,fill,stroke=None,round=False):
 s=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE,X(x),X(y),X(w),X(h))
 if round:s.adjustments[0]=.07
 s.fill.solid();s.fill.fore_color.rgb=C(fill)
 if stroke:s.line.color.rgb=C(stroke);s.line.width=Pt(1.3)
 else:s.line.fill.background()
 return s
def label(sl,t,x,y,w,h,size,col="182536",bold=False,right=False):
 s=sl.shapes.add_textbox(X(x),X(y),X(w),X(h))
 f=s.text_frame;f.margin_left=0;f.margin_top=0;f.margin_right=0;f.margin_bottom=0
 p=f.paragraphs[0];p.alignment=PP_ALIGN.RIGHT if right else PP_ALIGN.LEFT
 r=p.add_run();r.text=t;r.font.name="Arial";r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=C(col)
 return s
def heading(sl,n,sec,title):
 box(sl,0,0,1600,900,"F3F6F9")
 label(sl,sec,112,71,900,36,13.5,"52627B")
 label(sl,f"{n:02d} / 09",1330,71,160,36,13.5,"52627B",right=True)
 label(sl,title,112,135,1200,85,40.5,bold=True)
prs=Presentation(BASE)
assert len(prs.slides)==7
for sl in prs.slides:
 for sh in sl.shapes:
  if sh.has_text_frame:
   for p in sh.text_frame.paragraphs:
    for r in p.runs:
     r.text=re.sub(r'\b(0[1-7])\s*/\s*07\b',r'\1 / 09',r.text)
     if r.font.name and r.font.name not in ("Arial","Courier New"):r.font.name="Arial"

s=prs.slides.add_slide(prs.slide_layouts[0])
heading(s,8,"CONCLUSIONI E SVILUPPI FUTURI","Futuri sviluppi")
box(s,112,315,668,310,"F0F4FF","D9E3EF",True)
box(s,820,315,668,310,"EEF8F3","D9E3EF",True)
label(s,"Ampliare i test supportati",144,350,610,55,24.5,"4264A8",True)
label(s,"Automatizzare il mutation testing",852,350,610,55,24.5,"278766",True)
for x,items in [(144,["Estendere i normalizzatori a nuovi protocolli","e ricavare nuove evidenze, per coprire","più requisiti e casi di test SCAS."]),(852,["Utilizzare un unico motore comune","anche per i casi di mutation testing,","guidato dalle regole della DSL."])]:
 for j,v in enumerate(items):label(s,v,x,432+j*42,610,40,20.25,"52627B")

s=prs.slides.add_slide(prs.slide_layouts[0])
box(s,0,0,1600,900,"F3F6F9")
label(s,"09 / 09",1330,71,160,36,13.5,"52627B",right=True)
line=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,X(144),X(348),X(270),X(348))
line.line.color.rgb=C("60758F");line.line.width=Pt(3.75)
label(s,"Grazie per l’attenzione.",144,375,850,110,52.5,"182536",True)
# Right-hand original illustration, copied from the 20-minute closing slide.
original=(OUT.parent/"5g-thesis-presentation"/"slides"/"21-conclusioni.html").read_text(encoding="utf-8")
svg=re.search(r'<svg[\s\S]*?</svg>',original).group(0)
svg=re.sub(r'<text[^>]*>\s*Domande\?\s*</text>','',svg)
svg=svg.replace('viewBox="0 0 1600 900"','viewBox="1018 300 500 430"').replace('width="1600" height="900"','width="1000" height="860"')
# Render only the closing illustration; editable 'Grazie' text is native Arial.
png=cairosvg.svg2png(bytestring=svg.encode("utf-8"),output_width=1000,output_height=860)
s.shapes.add_picture(io.BytesIO(png),X(1018),X(300),X(500),X(430))
animated=OUT/"tesi-5min-9slide-animato.pptx"
prs.save(animated)
subprocess.run(["libreoffice","-env:UserInstallation=file:///tmp/lo_5g9_a","--headless","--convert-to","pdf","--outdir",str(OUT),str(animated)],check=True)
first=OUT/"tesi-5min-9slide-animato.pdf"
for i in (4,7):
 subprocess.run(["pdftoppm","-f",str(i),"-l",str(i),"-scale-to","1600","-png","-singlefile",str(first),str(OUT/f"last-{i}")],check=True)
static=Presentation(animated)
for i in (4,7):
 sl=static.slides[i-1]
 for shape in list(sl.shapes):
  el=shape._element;el.getparent().remove(el)
 sl.shapes.add_picture(str(OUT/f"last-{i}.png"),0,0,width=static.slide_width,height=static.slide_height)
for sl in static.slides:
 for el in sl._element.xpath("./p:timing"):el.getparent().remove(el)
st=OUT/"tesi-5min-9slide-statico.pptx"
static.save(st)
subprocess.run(["libreoffice","-env:UserInstallation=file:///tmp/lo_5g9_s","--headless","--convert-to","pdf","--outdir",str(OUT),str(st)],check=True)
(OUT/"tesi-5min-9slide-statico.pdf").replace(OUT/"tesi-5min-9slide.pdf")
with ZipFile(OUT/"presentazione-tesi-5min-9slide.zip","w",ZIP_DEFLATED,compresslevel=8) as z:
 for fn in ["index.html","tesi-5min-9slide-animato.pptx","tesi-5min-9slide-statico.pptx","tesi-5min-9slide.pdf"]:
  z.write(OUT/fn,arcname=fn)
print("Created 9-slide assets")
